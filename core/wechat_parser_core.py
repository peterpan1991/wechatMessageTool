import os
import re
from datetime import datetime
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


class WechatParserCore:
    def __init__(self):
        pass
    
    @staticmethod
    def get_wechat_user(navbar_path):
        with open(navbar_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        soup = BeautifulSoup(content, 'html.parser')
        
        wechat_user = None
        
        for a in soup.find_all('a', class_='jstree-anchor'):
            if '微信' in a.get_text():
                ul = a.find_next_sibling('ul')
                if ul:
                    for child_a in ul.find_all('a', class_='jstree-anchor'):
                        icon = child_a.find('i', class_='icon-userinfo')
                        if icon:
                            text = child_a.get_text()
                            text = ' '.join(text.split())
                            text = re.sub(r'[（\(]\d+/\d+[）\)]$', '', text)
                            wechat_user = text.strip()
                            return wechat_user
        
        return wechat_user
    
    @staticmethod
    def find_water_record_links(navbar_path):
        with open(navbar_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        soup = BeautifulSoup(content, 'html.parser')
        
        water_records = []
        target_a = None
        
        for a in soup.find_all('a', class_='jstree-anchor'):
            text = a.get_text()
            if '流水记录' in text and target_a is None:
                target_a = a
                break
        
        if target_a:
            ul = target_a.find_next_sibling('ul')
            if ul:
                for child_a in ul.find_all('a', class_='jstree-anchor'):
                    href = child_a.get('href')
                    if href:
                        parts = href.split('#')
                        page = parts[0]
                        record_id = parts[1] if len(parts) > 1 else ''
                        
                        name_text = child_a.get_text()
                        name_text = ' '.join(name_text.split())
                        name_text = re.sub(r'[（\(]\d+/\d+[）\)]$', '', name_text)
                        name = name_text.strip()
                        
                        water_records.append({
                            'name': name,
                            'page': page,
                            'id': record_id
                        })
        
        return water_records
    
    @staticmethod
    def parse_transaction_page(page_path, record_id=None):
        with open(page_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        soup = BeautifulSoup(content, 'html.parser')
        
        records = []

        target_section = None
        if record_id:
            target_h5 = soup.find('h5', id=record_id)
            if target_h5:
                next_div = target_h5.find_next_sibling('div')
                if next_div:
                    target_section = next_div.find('div', class_='conversationBox')
        
        if target_section:
            left_boxes = target_section.select('ul > li > div.leftBox, ul > li > div.rightBox')
        else:
            return records

        for left_box in left_boxes:
            try:
                name_elem = left_box.find('p', class_='ellipsis')
                if name_elem:
                    title = name_elem.get('title', '')
                    name_match = re.match(r'([^(]+)', title)
                    name = name_match.group(1).strip() if name_match else ''
                else:
                    name = ''
                
                date_span = left_box.find('span', style=re.compile(r'margin-left:\s*15px'))
                date = date_span.get_text(strip=True) if date_span else ''
                
                amount_div = left_box.find('div', class_='msg-red-amount')
                amount = amount_div.get_text(strip=True) if amount_div else ''
                
                type_div = left_box.find('div', class_='msg-red-name')
                category = type_div.get_text(strip=True) if type_div else ''
                
                trade_elem = left_box.find('p', class_='trade-number')
                trade_number = ''
                if trade_elem:
                    trade_text = trade_elem.get_text(strip=True)
                    trade_match = re.search(r'交易单号[：:]\s*(\S+)', trade_text)
                    if trade_match:
                        trade_number = trade_match.group(1)
                
                note_div = left_box.find('div', class_='msg-red-note')
                note = note_div.get_text(strip=True) if note_div else ''
                
                if not amount and not note:
                    title_elem = left_box.find('p', class_='msg-article-title')
                    summary_elem = left_box.find('p', class_='msg-article-summary')
                    if title_elem and summary_elem:
                        category = title_elem.get_text(strip=True)
                        summary_text = summary_elem.get_text(strip=True)
                        
                        amount_match = re.search(r'(?:付款|提现|收款|退还|收到|支付|红包|转账)金额[：]?[￥¥]?([\d.]+)', summary_text)
                        amount = f"￥{amount_match.group(1)}" if amount_match else ''
                        
                        trade_match = re.search(r'交易单号[：:](\S+)', summary_text)
                        if trade_match:
                            trade_number = trade_match.group(1)
                        
                        note = summary_text
                
                if name or date or amount or category:
                    records.append({
                        'chat_name': '',
                        'name': name,
                        'amount': amount,
                        'date': date,
                        'category': category,
                        'trade_number': trade_number,
                        'note': note
                    })
            except Exception as e:
                continue
        
        return records
    
    @staticmethod
    def create_excel(records, output_path):
        wb = Workbook()
        ws = wb.active
        ws.title = "微信流水记录"
        
        headers = ['对方名称', '发送方', '金额', '日期', '类别', '交易单号', '备注']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.fill = PatternFill(start_color='CCE5FF', end_color='CCE5FF', fill_type='solid')
        
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for row_idx, record in enumerate(records, 2):
            ws.cell(row=row_idx, column=1, value=record.get('chat_name', ''))
            ws.cell(row=row_idx, column=2, value=record.get('name', ''))
            ws.cell(row=row_idx, column=3, value=record.get('amount', ''))
            ws.cell(row=row_idx, column=4, value=record.get('date', ''))
            ws.cell(row=row_idx, column=5, value=record.get('category', ''))
            ws.cell(row=row_idx, column=6, value=record.get('trade_number', ''))
            ws.cell(row=row_idx, column=7, value=record.get('note', ''))
            
            for col in range(1, 8):
                ws.cell(row=row_idx, column=col).border = thin_border
        
        for col in range(1, 8):
            ws.column_dimensions[chr(64 + col)].width = 20
        
        wb.save(output_path)
        print(f"Excel file saved to: {output_path}")
    
    @staticmethod
    def process_folder(folder_path, progress_callback=None):
        folder_path = os.path.abspath(folder_path)
        
        navbar_path = os.path.join(folder_path, 'navbar.html')
        if not os.path.exists(navbar_path):
            return False, "navbar.html not found"
        
        if progress_callback:
            progress_callback(f"正在处理: {folder_path}", 0)
        
        core = WechatParserCore()
        wechat_user = core.get_wechat_user(navbar_path)
        
        water_records = core.find_water_record_links(navbar_path)
        
        all_records = []
        total = len(water_records)
        
        for idx, item in enumerate(water_records):
            page_path = os.path.join(folder_path, item['page'])
            if os.path.exists(page_path):
                if progress_callback:
                    progress_callback(f"正在解析 {item['page']} ({idx+1}/{total})...", int((idx + 1) / total * 100))
                records = core.parse_transaction_page(page_path, item['id'])
                for record in records:
                    record['chat_name'] = item['name']
                all_records.extend(records)
                if progress_callback:
                    progress_callback(f"Parsing {item['page']} (id: {item['id']}, name: {item['name'][:30]}... Found {len(records)} records", int((idx + 1) / total * 100))
        
        return True, (all_records, wechat_user)
